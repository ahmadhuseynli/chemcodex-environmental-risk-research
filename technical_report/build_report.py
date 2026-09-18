from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph


ROOT = Path(__file__).resolve().parent
SRC = ROOT / "TECHNICAL_NOTE_SOURCE.md"
DOCX = ROOT / "CHEMCODEX_ENVIRONMENTAL_RISK_TECHNICAL_NOTE_v0.1.docx"
PDF = ROOT / "CHEMCODEX_ENVIRONMENTAL_RISK_TECHNICAL_NOTE_v0.1.pdf"
SITE_PDF = ROOT.parent / "docs" / "assets" / PDF.name


def clean(text: str) -> str:
    return text.replace("**", "")


def parse_blocks():
    out = []
    for raw in SRC.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            out.append(("title", clean(line[2:])))
        elif line.startswith("## "):
            out.append(("heading", clean(line[3:])))
        elif line.startswith("- "):
            out.append(("bullet", clean(line[2:])))
        else:
            out.append(("body", clean(line)))
    return out


items = parse_blocks()

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)

heading = doc.styles["Heading 1"]
heading.font.name = "Arial"
heading._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
heading.font.size = Pt(15)
heading.font.bold = True
heading.font.color.rgb = RGBColor(39, 98, 74)

first = True
for kind, value in items:
    if kind == "title" and first:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(21)
        r.font.color.rgb = RGBColor(39, 98, 74)
        first = False
    elif kind == "heading":
        doc.add_heading(value, level=1)
    elif kind == "bullet":
        doc.add_paragraph(value, style="List Bullet")
    else:
        doc.add_paragraph(value)

for section in doc.sections:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ChemCodex / Environmental Chemical Hazard & Risk Research")
    r.font.name = "Arial"
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(105, 105, 105)

doc.core_properties.title = "ChemCodex - Environmental Chemical Hazard & Risk Research"
doc.core_properties.author = "Ahmad Huseynli"
doc.save(DOCX)

base = getSampleStyleSheet()
body = ParagraphStyle("Body", parent=base["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, spaceAfter=6)
h1 = ParagraphStyle("H1", parent=base["Heading1"], fontName="Helvetica-Bold", textColor=colors.HexColor("#27624A"), fontSize=15, leading=18, spaceBefore=10, spaceAfter=6)
title = ParagraphStyle("Title", parent=base["Title"], fontName="Helvetica-Bold", textColor=colors.HexColor("#27624A"), fontSize=21, leading=25, alignment=TA_CENTER, spaceAfter=14)

story = []
first = True
for kind, value in items:
    if kind == "title" and first:
        story.append(Paragraph(escape(value), title))
        first = False
    elif kind == "heading":
        story.append(Paragraph(escape(value), h1))
    elif kind == "bullet":
        story.append(Paragraph("&bull; " + escape(value), body))
    else:
        story.append(Paragraph(escape(value), body))

def footer(canvas, doc_obj):
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawString(18 * mm, 11 * mm, "ChemCodex public technical note")
    canvas.drawRightString(192 * mm, 11 * mm, str(doc_obj.page))
    canvas.restoreState()

pdf = SimpleDocTemplate(str(PDF), pagesize=A4, rightMargin=18*mm, leftMargin=18*mm, topMargin=16*mm, bottomMargin=18*mm)
pdf.build(story, onFirstPage=footer, onLaterPages=footer)

SITE_PDF.parent.mkdir(parents=True, exist_ok=True)
SITE_PDF.write_bytes(PDF.read_bytes())

print(DOCX)
print(PDF)
print(SITE_PDF)
